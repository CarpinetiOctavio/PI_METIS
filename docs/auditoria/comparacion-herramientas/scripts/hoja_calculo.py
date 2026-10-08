"""Las mismas 9 series en una hoja de cálculo con funciones nativas (las que
usaría un usuario de Excel), recalculada con LibreOffice Calc en modo headless.

Arma un .xlsx con fórmulas (no con valores), lo recalcula con
`soffice --headless --convert-to xlsx` y lee los resultados. Las funciones
usadas (PROMEDIO/AVERAGE, DESVEST.M/STDEV.S, COEFICIENTE.ASIMETRIA/SKEW,
CURTOSIS/KURT, COEF.DE.CORREL/CORREL, PRUEBA.T/T.TEST, DISTR.NORM.INV/NORM.INV,
DISTR.LOGNORM.INV/LOGNORM.INV, DISTR.GAMMA.INV/GAMMA.INV) tienen la misma
definición en Excel y en LibreOffice. Las de Excel 2010 en adelante se
escriben con el prefijo _xlfn., que es como Excel las guarda en el .xlsx.

    python docs/auditoria/comparacion-herramientas/scripts/hoja_calculo.py
"""

import csv
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
RES = AQUI.parent / "resultados"
PERIODOS = [2, 10, 100]


def armar(series: dict, ruta: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for clave, x in series.items():
        ws = wb.create_sheet(clave)
        n = len(x)
        ws["A1"] = "Q"
        for i, v in enumerate(x):
            ws.cell(row=2 + i, column=1, value=v)
        rng = f"A2:A{n + 1}"
        h = n // 2
        formulas = {
            "media": f"=AVERAGE({rng})",
            "desvio_s": f"=_xlfn.STDEV.S({rng})",
            "asimetria_skew": f"=SKEW({rng})",
            "curtosis_kurt": f"=KURT({rng})",
            # Autocorrelación de lag 1 "a lo Excel": CORREL de la serie contra
            # sí misma corrida un lugar (cada tramo con su propia media).
            "correl_lag1": f"=CORREL(A2:A{n},A3:A{n + 1})",
            # t de Student de mitades, varianzas iguales (tipo 2), dos colas.
            "ttest_p": f"=_xlfn.T.TEST(A2:A{h + 1},A{h + 2}:A{n + 1},2,2)",
            "mu_y": f"=SUMPRODUCT(LN({rng}))/{n}",
            "sigma_y": f"=SQRT(SUMPRODUCT((LN({rng})-SUMPRODUCT(LN({rng}))/{n})^2)/{n - 1})",
        }
        fila = 1
        for nombre, formula in formulas.items():
            ws.cell(row=fila, column=3, value=nombre)
            ws.cell(row=fila, column=4, value=formula)
            fila += 1
        # Cuantiles con las inversas nativas, parámetros por momentos (tesis
        # IV-92/93, IV-107/108, IV-123/124 y Gumbel IV-177/178).
        for t in PERIODOS:
            p = 1 - 1 / t
            cuant = {
                f"normal_T{t}": f"=_xlfn.NORM.INV({p},D1,D2)",
                f"lognormal2p_T{t}": f"=_xlfn.LOGNORM.INV({p},D7,D8)",
                f"gamma2p_T{t}": f"=_xlfn.GAMMA.INV({p},(D1/D2)^2,D2^2/D1)",
                f"gumbel_T{t}": f"=(D1-0.45*D2)-0.78*D2*LN(-LN({p}))",
            }
            for nombre, formula in cuant.items():
                ws.cell(row=fila, column=3, value=nombre)
                ws.cell(row=fila, column=4, value=formula)
                fila += 1
    wb.save(ruta)


def recalcular(ruta: Path, destino: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("No se encontró LibreOffice (soffice) en el PATH.")
    subprocess.run(
        [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(destino), str(ruta)],
        check=True,
        capture_output=True,
    )
    return destino / ruta.name


def main() -> None:
    series = defaultdict(list)
    with open(RES / "series.csv", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            series[r["estacion"]].append(float(r["valor"]))
    series = {k: v for k, v in series.items() if len(v) >= 10}
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        origen = tmp / "series_formulas.xlsx"
        armar(series, origen)
        (tmp / "out").mkdir()
        calc = recalcular(origen, tmp / "out")
        wb = openpyxl.load_workbook(calc, data_only=True)
        filas = []
        for clave in series:
            ws = wb[clave]
            for fila in range(1, ws.max_row + 1):
                nombre = ws.cell(row=fila, column=3).value
                if nombre:
                    filas.append({"estacion": clave, "magnitud": nombre, "hoja_calculo": ws.cell(row=fila, column=4).value})
        shutil.copy(origen, RES / "hoja_calculo_formulas.xlsx")
    with open(RES / "hoja_calculo.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["estacion", "magnitud", "hoja_calculo"], lineterminator="\n")
        w.writeheader()
        w.writerows(filas)
    version = subprocess.run(["soffice", "--version"], capture_output=True, text=True).stdout.strip()
    print(f"hoja_calculo.csv: {len(filas)} filas ({version})")


if __name__ == "__main__":
    main()
